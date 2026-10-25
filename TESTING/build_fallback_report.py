"""Build one plain Word report from a completed conditional fallback retest."""
from collections import Counter
from datetime import datetime, timezone
import argparse
import csv
import json
from pathlib import Path

from docx.shared import Pt
from build_report import document, percent, sanitize, table as base_table
from run_test import NAMES, metrics, load_labels

SHORT = {'DS1_Persian': 'Original print', 'DS1_Persian_Smoke': 'Modified print',
         'DS2_Persian_Hand': 'Original handwriting', 'DS2_Persian_Hand_Smoke': 'Modified handwriting'}


def table(doc, *args, **kwargs):
    result = base_table(doc, *args, **kwargs)
    spacer = doc.paragraphs[-1].paragraph_format
    spacer.line_spacing = Pt(4)
    spacer.space_before = spacer.space_after = Pt(0)
    return result


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def page(doc, title):
    heading = doc.add_heading(title, 1)
    heading.paragraph_format.page_break_before = True


def build(folder, output=None):
    folder = Path(folder).resolve()
    reports = [json.loads((folder / name / 'report.json').read_text(encoding='utf-8')) for name in NAMES]
    info = json.loads((folder / 'run_info.json').read_text(encoding='utf-8'))
    if info['status'] != 'complete' or any(r['run_status'] != 'complete' for r in reports):
        raise ValueError('All four datasets must be complete without technical errors.')
    labels = load_labels(Path(__file__).with_name('labels.json'))
    total = sum(r['overall']['images'] for r in reports)
    gained = sum(r['gained_correct'] for r in reports)
    wrong = sum(r['added_wrong_accepts'] for r in reports)
    doc = document('Conditional Fallback Retest Report',
        f'The TESTING batch runner evaluated all {total:,} images using the supplied YOLO checkpoint. '
        f'Conditional retries recovered {gained:,} additional correct predictions and accepted {wrong:,} '
        'additional incorrect predictions. Accepted BASE predictions and BASE multiple-letter rejections were preserved. '
        'Recovered predictions are flagged for review. The browser scanner and its interface were not changed.')
    doc.core_properties.author = 'Jeff Barrientos'
    doc.core_properties.last_modified_by = 'Jeff Barrientos'
    doc.core_properties.created = datetime.now(timezone.utc)
    doc.core_properties.modified = datetime.now(timezone.utc)
    doc.add_heading('Accuracy and accepted prediction reliability', 1)
    rows = []
    for r in reports:
        before, after = r['baseline'], r['overall']
        rows.append([SHORT[r['dataset']], percent(before['accuracy_percent']), percent(after['accuracy_percent']),
            percent(before['accepted_only_accuracy_percent']), percent(after['accepted_only_accuracy_percent'])])
    table(doc, ['Dataset', 'BASE accuracy', 'Fallback accuracy', 'BASE accepted correct', 'Fallback accepted correct'],
          rows, [2.0, 1.2, 1.2, 1.4, 1.4])
    doc.add_paragraph('Accuracy counts every image, including rejections. Accepted correct measures only the predictions '
                      'the program accepted. Recovering more letters can increase overall accuracy while reducing the '
                      'reliability of accepted answers.')
    doc.add_heading('Four recognition metrics after fallback', 1)
    table(doc, ['Dataset', 'Accuracy', 'Macro precision', 'Macro recall', 'Macro F1'],
        [[SHORT[r['dataset']], *[percent(r['overall'][k]) for k in
            ('accuracy_percent', 'macro_precision_percent', 'macro_recall_percent', 'macro_f1_percent')]] for r in reports],
        [2.0, 1.2, 1.4, 1.3, 1.3])
    doc.add_paragraph('Macro scores average the 32 letter classes equally. Rejected images count as misses. '
                      'These balanced datasets give matching accuracy and macro recall. These are image-level '
                      'recognition metrics, not bounding-box detection mAP.')
    doc.add_heading('Practical interpretation', 1)
    hand = reports[2]
    printed = reports[0]
    doc.add_paragraph(f"For original handwriting, accuracy changed from {percent(hand['baseline']['accuracy_percent'])} "
        f"to {percent(hand['overall']['accuracy_percent'])}, but {percent(hand['overall']['accepted_only_accuracy_percent'])} "
        'of accepted predictions were correct. Original printed macro F1 changed from '
        f"{percent(printed['baseline']['macro_f1_percent'])} to {percent(printed['overall']['macro_f1_percent'])}. "
        'Use the new recovered and needs_review columns to distinguish fallback '
        'answers. The results support further validation before adapting this logic into the browser scanner.')

    page(doc, 'Recognition rules and local verification')
    doc.add_paragraph('The first accepted stage supplies the final result. An accepted result is exactly one candidate '
        'at confidence 80% or higher. BASE multiple-letter rejections are final. Only zero-detection and low-confidence '
        'BASE outcomes enter the fallback. If all retries fail, the original BASE rejection is retained.')
    table(doc, ['Order', 'Method', 'Parameters'], [
        [1, 'BASE', 'Original decoded RGB image. No added enhancement.'],
        [2, 'Contrast', 'Grayscale; normalize the 90th percentile toward white; CLAHE clip 2.0 with a 4 by 4 tile grid.'],
        [3, 'Both', 'Contrast once, then 0, 45, 90, 135, 180, 225, 270 and 315 degrees counterclockwise.'],
        [4, 'Eight rotations', 'The same angles applied to the original image.'],
    ], [.5, 1.4, 5.3])
    doc.add_paragraph('Right-angle rotations use lossless transposes. Diagonals use bicubic interpolation on an expanded '
        'white canvas. At rotation stages, the strongest accepted prediction wins; equal scores use the lowest angle. '
        'Zero-degree results are reused, giving at most 16 model views per image. Confidence is a model score and is '
        'not a calibrated probability of correctness. No expected labels enter the selection logic.')
    controls_path = folder / 'negative_controls.csv'
    if controls_path.exists():
        doc.add_heading('Synthetic validation negative controls', 1)
        controls = read_csv(controls_path)
        table(doc, ['Input', 'Images', 'BASE false accepts', 'Fallback false accepts'],
            [[{'blank': 'Blank', 'multiple': 'Two letters'}.get(r['kind'], r['kind']),
              r['images'], r['base_false_accepts'], r['fallback_false_accepts']] for r in controls],
            [2.4, 1.0, 1.9, 1.9])
        doc.add_paragraph('Controls are excluded from the 32-class recognition metrics. These small synthetic controls '
                          'diagnose rejection behavior. They do not establish rejection '
                          'performance for real camera images, words, or natural handwriting.')
    doc.add_heading('Environment and evidence', 1)
    doc.add_paragraph(f"{info['device_name']}; FP32; batch {info['batch_size']}; PyTorch {info['torch']}; "
        f"Ultralytics {info['ultralytics']}. Input size 640 by 640, candidate floor 3%, NMS IoU 0.40, "
        'class-agnostic NMS, maximum 32 detections. The fixed acceptance threshold remains 80%.')
    doc.add_paragraph('Checkpoint, labels and code checksums are in run_info.json; input checksums are in the prediction CSVs. Each unique detector view '
        'was inferred fresh; identical visible pixels were reused only within this retest. Source images, weights '
        'and thresholds were unchanged. No training or refitting was performed.')
    doc.add_paragraph('Sources: github.com/mehr32/Image-dataset-of-Persian-letters (print) and '
        'github.com/myousefnezhad/persianhandwriting (hand). Smoke folders contain one previous augmentation '
        'per original; its parameters are retained in the prediction CSVs.')
    doc.add_paragraph('Flipped samples retain their original source labels, even when mirroring changes the visible '
        'glyph. Scores for those groups use the original-label assumption.')
    doc.add_paragraph('The fallback order was selected after inspecting earlier results from these same datasets. '
        'This retest verifies the implementation on those datasets and is not an independent estimate of generalization. '
        'No writer-grouped handwriting validation or manual annotation of every image was performed.')

    for r in reports:
        page(doc, SHORT[r['dataset']] + ' results')
        before, after = r['baseline'], r['overall']
        doc.add_paragraph(f"{after['images']:,} images; {after['technical_errors']} technical errors. "
            f"Additional correct predictions: {r['gained_correct']:,}. Additional wrong accepted predictions: "
            f"{r['added_wrong_accepts']:,}. Lost correct BASE predictions: {r['lost_correct']:,}.")
        table(doc, ['Method', 'Accuracy', 'Macro precision', 'Macro recall', 'Macro F1'],
            [[name, *[percent(s[k]) for k in ('accuracy_percent', 'macro_precision_percent', 'macro_recall_percent', 'macro_f1_percent')]]
             for name, s in [('BASE', before), ('Fallback', after)]], [1.4, 1.3, 1.6, 1.5, 1.4])
        doc.add_heading('Accepted and rejected outcomes', 1)
        def outcomes(s):
            n = s['images']
            c = s['status_counts']
            return {'Accepted correct': s['correct'], 'Accepted incorrect': c.get('accepted', 0) - s['correct'],
                'Low confidence': c.get('low_confidence', 0), 'No detection': c.get('no_detection', 0),
                'Multiple detections': c.get('multiple_detections', 0), 'Technical errors': s['technical_errors']}
        b, a = outcomes(before), outcomes(after)
        table(doc, ['Outcome', 'BASE images', 'BASE share', 'Fallback images', 'Fallback share'],
            [[name, f'{b[name]:,}', percent(100*b[name]/after['images']), f'{a[name]:,}', percent(100*a[name]/after['images'])] for name in b],
            [2.4, 1.1, 1.2, 1.3, 1.2])
        doc.add_heading('Method supplying accepted answers', 1)
        raw = read_csv(folder / r['dataset'] / 'predictions.csv')
        stage_rows = []
        for method in ('base', 'contrast', 'both', 'rotations'):
            subset = [row for row in raw if row['selected_method'] == method and row['status'] == 'accepted']
            count = len(subset)
            correct = sum(row['correct'] == 'yes' for row in subset)
            stage_rows.append([method, f'{count:,}', f'{correct:,}', f'{count-correct:,}', percent(100*correct/count) if count else 'No accepted answers'])
        table(doc, ['Method', 'Accepted', 'Correct', 'Incorrect', 'Accepted correct'], stage_rows, [1.4, 1.2, 1.2, 1.2, 2.2])
        doc.add_paragraph('BASE rows above include accepted answers only. Exhausted fallback attempts retain the BASE '
            'rejection and appear in the outcome table. All accepted fallback rows have needs_review=yes. '
            'No accepted BASE answer or BASE multiple-letter rejection was replaced.')
        if r['modified']:
            doc.add_heading('Original modification groups', 1)
            rows = []
            for kind in ('rotation', 'flip', 'rotationFlip', 'blur', 'brightnessContrast'):
                group = [row for row in raw if row['modification'] == kind]
                if not group:
                    continue
                count = len(group)
                rows.append([kind, f'{count:,}', percent(100*sum(x['base_correct'] == 'yes' for x in group)/count),
                    percent(100*sum(x['correct'] == 'yes' for x in group)/count)])
            table(doc, ['Modification', 'Images', 'BASE accuracy', 'Fallback accuracy'], rows, [2.5, 1.1, 1.8, 1.8])
    path = Path(output) if output else folder / 'Reports.docx'
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    sanitize(path)
    print('REPORT SAVED ' + str(path), flush=True)
    return path


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('folder', type=Path)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    build(args.folder, args.output)
