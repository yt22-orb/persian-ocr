"""Build one merged Word report with interpretation, four datasets, and error analysis."""
from pathlib import Path
import argparse
import csv
import json
import zipfile
from copy import deepcopy
from datetime import datetime, timezone

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ORDER=['DS1_Persian','DS1_Persian_Smoke','DS2_Persian_Hand','DS2_Persian_Hand_Smoke']
DISPLAY={'unmodified':'Unmodified','processed_all':'All modified','rotation':'Rotation','flip':'Flip',
         'rotationFlip':'Rotation and flip','blur':'Gaussian blur','brightnessContrast':'Brightness and contrast'}


def percent(value): return f'{value:.2f}%'


def sanitize(path):
    with zipfile.ZipFile(path) as archive:
        contents={item.filename:archive.read(item) for item in archive.infolist()}
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for name,raw in list(contents.items()):
        if not name.startswith('word/') or not name.endswith('.xml'): continue
        tree=etree.fromstring(raw)
        for tag in ('shd','highlight','background','pBdr'):
            for element in tree.findall('.//w:'+tag,ns): element.getparent().remove(element)
        for element in tree.findall('.//w:color',ns):
            element.set(qn('w:val'),'000000')
            for attr in list(element.attrib):
                if attr!=qn('w:val'): del element.attrib[attr]
        for element in tree.iter():
            if qn('w:color') in element.attrib: element.set(qn('w:color'),'000000')
        contents[name]=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    assert not any(name.startswith('word/media/') for name in contents)
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as archive:
        for name,raw in contents.items(): archive.writestr(name,raw)


def document(title,intro):
    doc=Document()
    section=doc.sections[0]
    section.page_width,section.page_height=Inches(8.5),Inches(11)
    section.top_margin=section.bottom_margin=Inches(.6)
    section.left_margin=section.right_margin=Inches(.65)
    for style in doc.styles:
        if hasattr(style,'font'):
            style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal']
    normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.line_spacing=1.05
    doc.styles['Title'].font.size=Pt(16)
    doc.styles['Title'].font.bold=False
    doc.styles['Title'].paragraph_format.space_after=Pt(7)
    doc.styles['Heading 1'].font.size=Pt(12)
    doc.styles['Heading 1'].font.bold=True
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
    doc.add_paragraph(title,'Title')
    doc.add_paragraph(intro)
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('Page '); run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
    footer._p.append(field)
    return doc


def table(doc,headers,rows,widths,font_size=10):
    item=doc.add_table(rows=1,cols=len(headers))
    item.alignment=WD_TABLE_ALIGNMENT.CENTER
    item.autofit=False
    for index,width in enumerate(widths): item.columns[index].width=Inches(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        child=OxmlElement('w:'+edge); child.set(qn('w:val'),'single')
        child.set(qn('w:sz'),'4'); child.set(qn('w:color'),'000000'); borders.append(child)
    item._tbl.tblPr.append(borders)
    for index,text in enumerate(headers): item.rows[0].cells[index].text=str(text)
    repeat=OxmlElement('w:tblHeader'); item.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=item.add_row().cells
        for index,value in enumerate(values): cells[index].text=str(value)
    for row_index,row in enumerate(item.rows):
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for col_index,cell in enumerate(row.cells):
            cell.width=Inches(widths[col_index])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                margin=OxmlElement('w:'+side); margin.set(qn('w:w'),'70'); margin.set(qn('w:type'),'dxa'); margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2)
                paragraph.paragraph_format.space_before=Pt(2)
                narrative=row_index>0 and len(cell.text.split())>12
                paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT if col_index==0 or narrative else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name='Arial'; run.font.size=Pt(font_size)
                    run.font.color.rgb=RGBColor(0,0,0); run.font.bold=row_index==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return item


def score_rows(conditions):
    return [[DISPLAY.get(row['condition'],row['condition']),f"{row['images']:,}",
             percent(row['accuracy_percent']),percent(row['macro_precision_percent']),
             percent(row['macro_recall_percent']),percent(row['macro_f1_percent'])] for row in conditions]


def dataset_document(report):
    name=report['dataset']; score=report['overall']; statuses=score['status_counts']; count=score['images']
    intro=(f"The supplied model correctly recognized {score['correct']:,} of {count:,} images "
           f"({percent(score['accuracy_percent'])}) in {name}. Macro F1 was {percent(score['macro_f1_percent'])}. "
           f"The test used the existing recognition policy, including the 80% confidence requirement.")
    doc=document(name.replace('_',' ')+' Model Test Report',intro)
    doc.add_heading('Test configuration',1)
    doc.add_paragraph(('Modified' if report['modified'] else 'Unmodified')+' images; supplied yolo11n.pt '
                      '(SHA256 '+report['model']['checkpoint_sha256'][:16]+'...). '+
                      report['model']['device_name']+'; FP32; 640 pixel input. '
                      'Candidate floor 3%; NMS IoU 0.40; accepted only with one candidate at confidence 80% or higher.')
    doc.add_heading('Recognition results',1)
    table(doc,['Condition','Images','Accuracy','Macro precision','Macro recall','Macro F1'],score_rows(report['conditions']),[1.65,.65,1.05,1.3,1.25,1.3])
    doc.add_heading('Recognition outcomes',1)
    accepted=statuses.get('accepted',0)
    outcome_rows=[['Accepted and correct',score['correct'],percent(score['accuracy_percent'])],
                  ['Accepted but incorrect',accepted-score['correct'],percent(100*(accepted-score['correct'])/count)],
                  ['Low confidence',statuses.get('low_confidence',0),percent(100*statuses.get('low_confidence',0)/count)],
                  ['No detection',statuses.get('no_detection',0),percent(100*statuses.get('no_detection',0)/count)],
                  ['Multiple detections',statuses.get('multiple_detections',0),percent(100*statuses.get('multiple_detections',0)/count)],
                  ['Technical errors',score['technical_errors'],percent(100*score['technical_errors']/count)]]
    outcome_rows=[[row[0],f'{row[1]:,}',row[2]] for row in outcome_rows]
    table(doc,['Outcome','Images','Share of all images'],outcome_rows,[3.65,1.3,2.25])
    doc.add_heading('Interpretation',1)
    doc.add_paragraph(f"The recognizer accepted {percent(score['accepted_percent'])} of images. Among accepted predictions, "
                      f"{percent(score['accepted_only_accuracy_percent'])} were correct; this diagnostic excludes rejected images, "
                      'so the overall accuracy above remains the main result.')
    doc.add_paragraph('Macro scores average the 32 letter scores equally. Rejections count as misses for the expected letter. '
                      'Training independence has not been established; these scores describe the supplied folders under the saved policy.')
    doc.add_paragraph('Full records: the corresponding CSV and JSON results. Sources and preparation: Test Set ReadMe.md.')
    return doc


def interpretation_document(reports,paired):
    by_name={report['dataset']:report for report in reports}
    printed_before=by_name['DS1_Persian']['overall']; printed_after=by_name['DS1_Persian_Smoke']['overall']
    hand_before=by_name['DS2_Persian_Hand']['overall']; hand_after=by_name['DS2_Persian_Hand_Smoke']['overall']
    intro=(f"Printed accuracy changed from {percent(printed_before['accuracy_percent'])} to {percent(printed_after['accuracy_percent'])} "
           f"after modification. Handwritten accuracy changed from {percent(hand_before['accuracy_percent'])} to "
           f"{percent(hand_after['accuracy_percent'])}. The model recognized a smaller share of handwritten images than printed images; rejected recognitions remain in all overall scores.")
    doc=document('Persian Letter Recognition Test Reports',intro)
    doc.add_heading('Four test set results',1)
    rows=[]
    short_names={'DS1_Persian':'DS1 printed original','DS1_Persian_Smoke':'DS1 printed modified',
                 'DS2_Persian_Hand':'DS2 hand original','DS2_Persian_Hand_Smoke':'DS2 hand modified'}
    for report in reports:
        score=report['overall']
        rows.append([short_names[report['dataset']],f"{score['images']:,}",percent(score['accuracy_percent']),
                     percent(score['macro_precision_percent']),percent(score['macro_recall_percent']),percent(score['macro_f1_percent'])])
    table(doc,['Dataset','Images','Accuracy','Macro precision','Macro recall','Macro F1'],rows,[1.65,.65,1.05,1.3,1.25,1.3])
    doc.add_heading('What changed',1)
    comparisons=[]
    for domain,before,after in [('Printed',printed_before,printed_after),('Handwritten',hand_before,hand_after)]:
        delta=after['accuracy_percent']-before['accuracy_percent']
        fdelta=after['macro_f1_percent']-before['macro_f1_percent']
        meaning=f"Correct recognition {'fell' if delta<0 else 'rose'} by {abs(delta):.2f} percentage points. "
        meaning+=f"Modified inputs were rejected in {percent(after['rejection_percent'])} of cases."
        comparisons.append([domain,f'{delta:+.2f} pp',f'{fdelta:+.2f} pp',meaning])
    table(doc,['Comparison','Accuracy change','F1 change','Interpretation'],comparisons,[1.0,1.15,1.05,4.0])
    doc.add_heading('Reading the four metrics',1)
    table(doc,['Metric','Meaning in this test'],[
        ['Accuracy','The percentage of all images that received the correct accepted prediction.'],
        ['Macro precision','How often an accepted letter assignment is correct, averaged equally across letters.'],
        ['Macro recall','How many images of each letter are correctly recognized, averaged equally across letters.'],
        ['Macro F1','Balances precision and recall per letter, then averages those F1 scores equally.'],
    ],[1.5,5.7])
    doc.add_paragraph('Accuracy and macro recall match because the letter supports are balanced. High precision can coexist '
                      'with lower accuracy when many inputs are rejected. Use overall accuracy and macro F1 together; accepted-only accuracy describes only the accepted subset.')
    doc.add_page_break()
    doc.add_heading('Comparisons using the same original images',1)
    doc.add_paragraph('Each row compares a modification group with the unmodified versions of those exact images. '
                      'The script checks source paths, labels, and checksums before forming these pairs. A negative change means reduced accuracy.')
    for name,title in [('DS1_Persian_Smoke','Printed modification groups'),('DS2_Persian_Hand_Smoke','Handwritten modification groups')]:
        doc.add_heading(title,1)
        group=[row for row in paired if row['dataset']==name]
        order=['rotation','flip','rotationFlip','blur','brightnessContrast']
        group.sort(key=lambda row:order.index(row['modification']))
        rows=[[DISPLAY[row['modification']],f"{int(row['images']):,}",percent(float(row['unmodified_accuracy_percent'])),
               percent(float(row['modified_accuracy_percent'])),percent(float(row['unmodified_macro_f1_percent'])),
               percent(float(row['modified_macro_f1_percent'])),f"{float(row['accuracy_change_percentage_points']):+.2f}"] for row in group]
        table(doc,['Category','Images','Original accuracy','Modified accuracy','Original F1','Modified F1','Accuracy change pp'],rows,[1.6,.6,1.0,1.0,1.0,1.0,1.0],font_size=9.5)
        worst=min(group,key=lambda row:float(row['accuracy_change_percentage_points']))
        doc.add_paragraph(f"The largest accuracy decrease was {DISPLAY[worst['modification']].lower()}: "
                          f"{float(worst['accuracy_change_percentage_points']):+.2f} percentage points against its matched originals.")
    doc.add_heading('Scope of the conclusions',1)
    doc.add_paragraph('The model and its 80% acceptance threshold were fixed before testing. No training or threshold tuning was performed. '
                      'Independence from earlier training data has not been verified, so these results do not establish unseen-font or unseen-writer performance.')
    doc.add_paragraph('Sources and preparation details are documented in Test Set ReadMe.md. Per-image outputs and category metrics are in the four results folders. The following pages contain the four dataset reports.')
    return doc


def build(folder, output=None):
    folder=Path(folder).resolve()
    reports=[json.loads((folder/name/'report.json').read_text(encoding='utf-8')) for name in ORDER]
    if not all(report['run_status']=='complete' for report in reports):
        raise ValueError('All four dataset runs must be complete before building the merged report.')
    paired=list(csv.DictReader((folder/'paired_comparison.csv').open(encoding='utf-8-sig')))
    merged=interpretation_document(reports,paired)
    for report in reports:
        merged.add_page_break()
        section=dataset_document(report)
        for element in section.element.body:
            if element.tag != qn('w:sectPr'):
                merged.element.body.insert(len(merged.element.body)-1,deepcopy(element))
    from error_analysis import append_error_analysis
    append_error_analysis(merged,folder,table)
    merged.core_properties.author='Jeff Barrientos'
    merged.core_properties.last_modified_by='Jeff Barrientos'
    merged.core_properties.title='Persian Letter Recognition Test Reports'
    merged.core_properties.subject='Baseline and modified printed and handwritten test results'
    merged.core_properties.created=datetime.now(timezone.utc)
    merged.core_properties.modified=datetime.now(timezone.utc)
    target=Path(output) if output else folder/'Reports.docx'
    merged.save(target)
    sanitize(target)
    print('MERGED REPORT READY '+str(target),flush=True)
    return target


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('results',type=Path,nargs='?',default=Path(__file__).with_name('Results'))
    parser.add_argument('--output',type=Path)
    args=parser.parse_args()
    build(args.results,args.output)
